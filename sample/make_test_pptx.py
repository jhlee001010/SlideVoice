from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
layout = prs.slide_layouts[1]

slides_data = [
    ("첫 번째 슬라이드", "안녕하세요. 첫 번째 슬라이드에 대한 설명입니다."),
    ("두 번째 슬라이드", "이제 두 번째 슬라이드로 넘어가겠습니다."),
]

for title, note in slides_data:
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.notes_slide.notes_text_frame.text = note

prs.save("sample/test.pptx")
print("wrote sample/test.pptx")
